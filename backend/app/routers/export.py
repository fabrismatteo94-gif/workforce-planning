import io
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .. import models
from ..database import get_db

router = APIRouter(prefix="/aziende/{azienda_id}/export", tags=["export"])


@router.get("/excel")
def export_excel(azienda_id: str, db: Session = Depends(get_db)):
    azienda = db.query(models.Azienda).filter(models.Azienda.id == azienda_id).first()
    if not azienda:
        raise HTTPException(status_code=404, detail="Azienda non trovata")

    ultimo_snapshot = (
        db.query(models.CalcoloSnapshot)
        .filter(models.CalcoloSnapshot.azienda_id == azienda_id)
        .order_by(models.CalcoloSnapshot.data_calcolo.desc())
        .first()
    )
    if not ultimo_snapshot:
        raise HTTPException(status_code=400, detail="Nessun calcolo ancora eseguito per questa azienda")

    wb = Workbook()
    ws = wb.active
    ws.title = "Per Funzione"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="000000", end_color="000000", fill_type="solid")

    ws.append(["Funzione", "FTE calcolato", "Confidence", "Formula usata"])
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill

    for _, dati in ultimo_snapshot.risultati_per_funzione.items():
        ws.append([dati["nome"], dati["fte_finale"], dati["confidence"], dati["formula_usata"]])

    ws2 = wb.create_sheet("Riepilogo")
    ws2.append(["Azienda", azienda.nome])
    ws2.append(["Scenario", ultimo_snapshot.scenario_tipo])
    ws2.append(["FTE totali", ultimo_snapshot.fte_totali])
    ws2.append(["Costo totale annuo (€)", ultimo_snapshot.costo_totale_annuo])
    ws2.append(["Budget utilizzo (%)", ultimo_snapshot.budget_utilizzo_pct])
    ws2.append(["Data calcolo", ultimo_snapshot.data_calcolo.isoformat()])

    for col_ws in (ws, ws2):
        for column_cells in col_ws.columns:
            length = max(len(str(c.value)) if c.value is not None else 0 for c in column_cells)
            col_ws.column_dimensions[column_cells[0].column_letter].width = max(12, length + 2)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"workforce_plan_{azienda.nome.replace(' ', '_')}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
